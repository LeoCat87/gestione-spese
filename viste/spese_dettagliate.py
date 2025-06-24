import streamlit as st
import openpyxl
import pandas as pd

def mostra_spese_dettagliate(EXCEL_PATH, anno, persona, nome_foglio, carica_spese, carica_riepilogo, formatta_euro):
    st.title("📌 Spese Dettagliate")
    df_spese = carica_spese()
    df_riepilogo = carica_riepilogo()

    mesi_disponibili = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
                        "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]

    mappa_macrocategorie = {
        "Entrate": ["Stipendio", "Entrate extra", "Affitto Savoldo 4 + generico"],
        "Uscite necessarie": [
            "Affitto", "Bollette", "Spesa", "Abbonamenti", "Trasporti", "Assicurazione",
            "PAC Investimenti", "Mutuo", "Luce&Gas", "Internet/Telefono", "Mezzi",
            "Spese condominiali", "Spese comuni", "Auto (benzina, noleggio, pedaggi, parcheggi)",
            "Spesa cibo", "Tari", "Unobravo", "Donazioni (StC, Unicef, Greenpeace)"
        ],
        "Uscite variabili": [
            "Amazon", "Bolli governativi", "Farmacia/Visite", "Food Delivery", "Generiche",
            "Multa", "Uscite (Pranzi,Cena,Apericena,Pub,etc)", "Prelievi", "Regali",
            "Sharing (auto, motorino, bici)", "Shopping (vestiti, mobili,...)", "Stireria",
            "Viaggi (treno, aereo, hotel, attrazioni, concerti, cinema)"
        ]
    }

    col1, col2 = st.columns([1, 5])
    with col2:
        st.markdown("### ➕ Inserisci una nuova spesa")

        nuovo_testo = st.text_input("Descrizione", "")
        nuovo_valore = st.number_input("Importo (€)", step=0.01, format="%.2f")
        nuova_categoria = st.selectbox("Macrocategoria", list(mappa_macrocategorie.keys()))
        nuovo_tag = st.selectbox("Tag", mappa_macrocategorie[nuova_categoria])
        nuovo_mese = st.selectbox("Mese", mesi_disponibili)

        if st.button("➕ Aggiungi spesa"):
            if nuovo_testo.strip() == "" or nuovo_valore == 0:
                st.warning("⚠️ Inserisci una descrizione e un valore diverso da zero.")
            else:
                wb = openpyxl.load_workbook(EXCEL_PATH)
                ws = wb[nome_foglio("Spese")]

                mesi_excel = ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno",
                              "luglio", "agosto", "settembre", "ottobre", "novembre", "dicembre"]
                mese_col_start = None
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=1, column=col).value
                    if val and isinstance(val, str) and val.lower() == nuovo_mese.lower():
                        mese_col_start = col
                        break

                if mese_col_start:
                    row_idx = 3
                    while ws.cell(row=row_idx, column=mese_col_start).value not in [None, ""]:
                        row_idx += 1

                    ws.cell(row=row_idx, column=mese_col_start).value = nuovo_testo
                    ws.cell(row=row_idx, column=mese_col_start + 1).value = float(nuovo_valore)
                    ws.cell(row=row_idx, column=mese_col_start + 2).value = nuovo_tag

                    wb.save(EXCEL_PATH)
                    st.success("✅ Spesa aggiunta correttamente.")
                    st.experimental_rerun()
                else:
                    st.error("❌ Colonna del mese non trovata nel foglio Excel.")

    mesi_selezionati = st.multiselect("📅 Filtra per mesi:", mesi_disponibili, default=mesi_disponibili)
    categorie_tag = sorted([str(tag) for tag in df_riepilogo.index if pd.notnull(tag)])
    tag_selezionati = st.multiselect("🏷️ Filtra per categorie (Tag):", ["Tutti"] + categorie_tag, default=["Tutti"])

    df_filtrato = df_spese[df_spese["Mese"].isin(mesi_selezionati)].copy()
    if "Tutti" not in tag_selezionati:
        df_filtrato = df_filtrato[df_filtrato["Tag"].isin(tag_selezionati)]

    if not df_filtrato.empty:
        totale = df_filtrato["Valore"].sum()
        st.markdown(f"**Totale spese filtrate:** {formatta_euro(totale)}")
    else:
        st.info("🔍 Nessuna spesa trovata con i filtri selezionati.")

    if len(mesi_selezionati) == 1 and not df_filtrato.empty:
        edited_df = st.data_editor(
            df_filtrato[["Testo", "Valore", "Tag"]],
            use_container_width=False,
            hide_index=True,
            column_config={
                "Testo": st.column_config.TextColumn(label="Descrizione", help="Testo libero"),
                "Valore": st.column_config.NumberColumn(label="€", help="Importo della spesa", step=0.01, format="€ %.2f"),
                "Tag": st.column_config.SelectboxColumn(label="Tag", help="Categoria", options=categorie_tag, required=True)
            }
        )

        if not edited_df.equals(df_filtrato[["Testo", "Valore", "Tag"]]):
            st.success("✅ Modifiche rilevate.")
            if st.button("💾 Salva modifiche"):
                mese_sel = mesi_selezionati[0]
                df_aggiornato = df_spese[df_spese["Mese"] != mese_sel].copy()
                edited_df["Mese"] = mese_sel
                edited_df["Valore"] = pd.to_numeric(edited_df["Valore"], errors="coerce").fillna(0)

                def categoria_per_tag(tag):
                    if tag in mappa_macrocategorie["Entrate"]:
                        return "Entrate"
                    elif tag in mappa_macrocategorie["Uscite necessarie"]:
                        return "Uscite necessarie"
                    else:
                        return "Uscite variabili"

                edited_df["Categoria"] = edited_df["Tag"].apply(categoria_per_tag)
                edited_df["Testo"] = edited_df["Testo"].fillna("")

                df_finale = pd.concat([df_aggiornato, edited_df], ignore_index=True)

                wb = openpyxl.load_workbook(EXCEL_PATH)
                ws = wb[nome_foglio("Spese")]

                mese_col_start = None
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=1, column=col).value
                    if val and isinstance(val, str) and val.lower() == mese_sel.lower():
                        mese_col_start = col
                        break

                if mese_col_start:
                    for row in range(3, ws.max_row + 1):
                        for c in range(mese_col_start, mese_col_start + 3):
                            ws.cell(row=row, column=c).value = None

                    ws.cell(row=2, column=mese_col_start).value = "Testo"
                    ws.cell(row=2, column=mese_col_start + 1).value = "Valore"
                    ws.cell(row=2, column=mese_col_start + 2).value = "Tag"

                    for i, row in edited_df.iterrows():
                        ws.cell(row=3 + i, column=mese_col_start).value = row["Testo"]
                        ws.cell(row=3 + i, column=mese_col_start + 1).value = float(row["Valore"])
                        ws.cell(row=3 + i, column=mese_col_start + 2).value = row["Tag"]

                    wb.save(EXCEL_PATH)
                    st.success("✅ Modifiche salvate correttamente.")
                else:
                    st.error("❌ Colonna del mese non trovata nel foglio Excel.")
    elif len(mesi_selezionati) != 1:
        st.info("✏️ Per modificare le spese, seleziona **un solo mese**.")
    else:
        st.info("🔍 Nessuna spesa da modificare per i filtri attivi.")

    with open(EXCEL_PATH, "rb") as f:
        st.download_button(
            label="📥 Scarica file aggiornato",
            data=f,
            file_name="Spese_App_aggiornato.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
